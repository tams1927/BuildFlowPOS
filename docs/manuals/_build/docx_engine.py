"""Reusable DOCX building engine for the Phase UM-1 manuals.

Provides a professional document scaffold: cover page, auto-updating Table of
Contents, page headers/footers with page numbers, heading styles, numbered
steps, embedded figures with captions, and Note / Tip / Warning callouts.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

ASSETS = os.path.join(os.path.dirname(__file__), "..", "_assets")

NAVY   = RGBColor(0x1E, 0x40, 0xAF)
DARK   = RGBColor(0x0F, 0x17, 0x2A)
SLATE  = RGBColor(0x47, 0x55, 0x69)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

CALLOUT = {
    "note":    ("Note",    "DCE9FB", "1E40AF"),
    "tip":     ("Tip",     "E3F6E8", "1B7F3B"),
    "warning": ("Warning", "FDE6E6", "B42318"),
}


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def _field(paragraph, instr):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(t); run._r.append(s); run._r.append(e)
    return run


class Manual:
    def __init__(self, title, subtitle, audience, confidentiality):
        self.title = title
        self.subtitle = subtitle
        self.audience = audience
        self.confidentiality = confidentiality
        self.fig = 0
        self.doc = Document()
        self._setup_styles()
        self._setup_page()

    # ---------- setup ----------
    def _setup_styles(self):
        d = self.doc
        normal = d.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, color in (("Heading 1", 18, NAVY), ("Heading 2", 14, DARK), ("Heading 3", 12, SLATE)):
            st = d.styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.color.rgb = color
            st.font.bold = True

    def _setup_page(self):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, m, Inches(1))

    def _enable_update_fields(self):
        settings = self.doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    def _header_footer(self):
        sec = self.doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = f"Hardware Supply POS  \u2022  {self.title}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(8); r.font.color.rgb = SLATE
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run(f"{self.confidentiality}    |    Version 1.0    |    Page ")
        r1.font.size = Pt(8); r1.font.color.rgb = SLATE
        _field(fp, "PAGE")
        r2 = fp.add_run(" of ")
        r2.font.size = Pt(8); r2.font.color.rgb = SLATE
        _field(fp, "NUMPAGES")
        for r in fp.runs:
            r.font.size = Pt(8); r.font.color.rgb = SLATE

    # ---------- cover + toc ----------
    def cover(self):
        d = self.doc
        for _ in range(4):
            d.add_paragraph()
        bar = d.add_paragraph(); bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = bar.add_run("HARDWARE SUPPLY POS")
        r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = NAVY
        sub = d.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = sub.add_run("Cloud Point-of-Sale & Inventory Platform")
        rs.font.size = Pt(11); rs.font.color.rgb = SLATE
        d.add_paragraph()
        t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rt = t.add_run(self.title)
        rt.font.size = Pt(34); rt.font.bold = True; rt.font.color.rgb = DARK
        st = d.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rst = st.add_run(self.subtitle)
        rst.font.size = Pt(13); rst.font.color.rgb = SLATE
        for _ in range(6):
            d.add_paragraph()
        meta = d.add_table(rows=4, cols=2)
        meta.alignment = 1
        _no_borders(meta)
        rows = [("Version", "1.0"), ("Audience", self.audience),
                ("Classification", self.confidentiality), ("Release date", "Version 1.0 \u2013 General Availability")]
        for i, (k, v) in enumerate(rows):
            c0, c1 = meta.rows[i].cells
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rr = p0.add_run(k + "   "); rr.font.bold = True; rr.font.color.rgb = NAVY; rr.font.size = Pt(11)
            p1 = c1.paragraphs[0]
            rv = p1.add_run(v); rv.font.size = Pt(11)
        self._page_break()

    def toc(self):
        h = self.doc.add_paragraph("Table of Contents")
        h.style = self.doc.styles["Heading 1"]
        p = self.doc.add_paragraph()
        _field(p, 'TOC \\o "1-2" \\h \\z \\u')
        note = self.doc.add_paragraph()
        rn = note.add_run("If the contents above appear empty, right-click and choose \u201cUpdate Field\u201d.")
        rn.font.italic = True; rn.font.size = Pt(9); rn.font.color.rgb = SLATE
        self._page_break()

    # ---------- content blocks ----------
    def _page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text, num=None):
        self._page_break()
        label = f"{num}.  {text}" if num else text
        self.doc.add_paragraph(label, style=self.doc.styles["Heading 1"])

    def h2(self, text):
        self.doc.add_paragraph(text, style=self.doc.styles["Heading 2"])

    def p(self, text, bold=False):
        para = self.doc.add_paragraph()
        r = para.add_run(text); r.font.bold = bold
        return para

    def label_p(self, label, text):
        para = self.doc.add_paragraph()
        rl = para.add_run(f"{label}  "); rl.font.bold = True; rl.font.color.rgb = NAVY
        para.add_run(text)
        return para

    def steps(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Number")

    def bullets(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def callout(self, kind, text):
        label, fill, txt = CALLOUT[kind]
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = 1
        cell = table.rows[0].cells[0]
        _shade(cell, fill)
        para = cell.paragraphs[0]
        rl = para.add_run(f"{label}: ")
        rl.font.bold = True; rl.font.color.rgb = RGBColor.from_string(txt)
        rt = para.add_run(text)
        rt.font.color.rgb = RGBColor.from_string("1A1A1A")
        self.doc.add_paragraph()

    def figure(self, key, caption):
        path = os.path.join(ASSETS, key + ".png")
        if not os.path.exists(path):
            self.callout("warning", f"Screenshot '{key}' is unavailable in this build.")
            return
        self.fig += 1
        max_w, max_h = 6.3, 8.2
        with Image.open(path) as im:
            w, h = im.size
        width_in = max_w
        height_in = max_w * h / w
        if height_in > max_h:
            height_in = max_h
            width_in = max_h * w / h
        pic_p = self.doc.add_paragraph(); pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pic_p.add_run()
        run.add_picture(path, width=Inches(width_in))
        cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = cap.add_run(f"Figure {self.fig}: {caption}")
        rc.font.italic = True; rc.font.size = Pt(9); rc.font.color.rgb = SLATE

    # ---------- render a section dict ----------
    def section(self, sec, num):
        self.h1(sec["heading"], num)
        for block in sec["blocks"]:
            kind = block[0]
            if kind == "p":
                self.p(block[1])
            elif kind == "h2":
                self.h2(block[1])
            elif kind == "purpose":
                self.label_p("Purpose.", block[1])
            elif kind == "expected":
                self.label_p("Expected result.", block[1])
            elif kind == "steps":
                self.steps(block[1])
            elif kind == "bullets":
                self.bullets(block[1])
            elif kind in ("note", "tip", "warning"):
                self.callout(kind, block[1])
            elif kind == "fig":
                self.figure(block[1], block[2])

    def build(self, sections, out_path):
        self.cover()
        self._header_footer()
        self.toc()
        for i, sec in enumerate(sections, start=1):
            self.section(sec, i)
        self._enable_update_fields()
        self.doc.save(out_path)
        return out_path
